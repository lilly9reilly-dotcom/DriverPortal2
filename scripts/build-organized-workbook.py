#!/usr/bin/env python3
"""Build the organized agent workbook so the result is visible without Google login."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path(sys.argv[1] if len(sys.argv) > 1 else "/opt/cursor/artifacts/sheet-organization")
OUT_DIR.mkdir(parents=True, exist_ok=True)

DATA = json.loads(Path("/tmp/org-seed.json").read_text(encoding="utf-8"))
groups = DATA["groups"]
headers = DATA["headers"]

wb = Workbook()

thin = Border(
    left=Side(style="thin", color="D0D7E2"),
    right=Side(style="thin", color="D0D7E2"),
    top=Side(style="thin", color="D0D7E2"),
    bottom=Side(style="thin", color="D0D7E2"),
)
head_fill = PatternFill("solid", fgColor="1E4D93")
head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
title_font = Font(name="Calibri", bold=True, color="183A72", size=16)
cell_font = Font(name="Calibri", size=12)
agent_fills = [
    "E8F1FF",
    "E9F8F0",
    "FFF6E5",
    "F3E9FF",
    "FFECEC",
    "E7F7FB",
    "F4F1EA",
]


def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin


def autosize(ws, min_w=14, max_w=36):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_w
        for cell in col:
            width = max(width, min(max_w, len(str(cell.value or "")) + 4))
        ws.column_dimensions[letter].width = width
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"


# تنظيم_المعتمدين
ws = wb.active
ws.title = "تنظيم_المعتمدين"
ws["A1"] = "ترتيب الشيتات — كل معتمد وكل سيارة من 2026_06 حتى الآن"
ws["A1"].font = title_font
ws.merge_cells("A1:D1")
ws["A2"] = "هذه نتيجة التشغيل هنا. الشيت الحي يتغيّر بعد رفع السكربت وتشغيل organizeSheetsNow"
ws.merge_cells("A2:D2")
ws.append([])
ws.append(["المعتمد / الجهة", "النوع", "رقم السيارة", "ورقة القاعدة"])
style_header(ws, 4, 4)
row_i = 5
for idx, group in enumerate(groups):
    fill = PatternFill("solid", fgColor=agent_fills[idx % len(agent_fills)])
    for car in group["cars"]:
        ws.append([group["name"], group["kind"], car["carNumber"], group["dbSheet"]])
        for col in range(1, 5):
            cell = ws.cell(row=row_i, column=col)
            cell.fill = fill
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin
        row_i += 1
autosize(ws)
ws.freeze_panes = "A5"

# Agents
ws = wb.create_sheet("Agents")
ws.append(["الاسم", "النوع", "فعال", "ملاحظات"])
style_header(ws, 1, 4)
for group in groups:
    ws.append([group["name"], group["kind"], "1", "مسجّل من قائمة المعتمدين"])
autosize(ws)

# Fleet
ws = wb.create_sheet("Fleet")
ws.append(["رقم السيارة", "السائق الافتراضي", "تابع لـ", "اسم المعتمد", "فعال", "ملاحظات"])
style_header(ws, 1, 6)
for group in groups:
    for car in group["cars"]:
        ws.append([car["carNumber"], "", group["kind"], group["name"], "1", ""])
autosize(ws)

# DB sheets
for group in groups:
    ws = wb.create_sheet(group["dbSheet"][:31])
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.append(["بانتظار ترحيل وصولات 2026_06 إلى الشهر الحالي من الشيت الحي", "", "", "", "", "", "", "", "", "", "", "", "", "", group["name"], "", ""])
    autosize(ws)

ws = wb.create_sheet("DB_غير_مصنف")
ws.append(headers)
style_header(ws, 1, len(headers))
autosize(ws)

xlsx_path = OUT_DIR / "تنظيم_المعتمدين_والسيارات.xlsx"
wb.save(xlsx_path)

# HTML preview
rows_html = []
for idx, group in enumerate(groups):
    bg = "#" + agent_fills[idx % len(agent_fills)]
    for car in group["cars"]:
        rows_html.append(
            f"<tr style='background:{bg}'><td>{group['name']}</td><td>{group['kind']}</td>"
            f"<td class='car'>{car['carNumber']}</td><td>{group['dbSheet']}</td></tr>"
        )

tabs = ["تنظيم_المعتمدين", "Agents", "Fleet"] + [g["dbSheet"] for g in groups] + ["DB_غير_مصنف"]
tab_html = "".join(f"<button class='tab{' active' if i==0 else ''}'>{name}</button>" for i, name in enumerate(tabs))

html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
  <meta charset="UTF-8">
  <title>نتيجة ترتيب الشيتات</title>
  <link href="https://fonts.googleapis.com/css2?family=Tajawal:wght@500;800&display=swap" rel="stylesheet">
  <style>
    body {{ margin:0; font-family:Tajawal,sans-serif; background:#eef2f7; color:#17253f; }}
    .bar {{ background:#0b3a74; color:#fff; padding:16px 22px; }}
    .bar h1 {{ margin:0; font-size:26px; }}
    .bar p {{ margin:6px 0 0; opacity:.9; }}
    .tabs {{ display:flex; gap:6px; overflow:auto; padding:10px 16px; background:#fff; border-bottom:1px solid #d7deea; }}
    .tab {{ border:0; background:#e8eef8; color:#1a335e; border-radius:8px 8px 0 0; padding:8px 12px; font-weight:800; }}
    .tab.active {{ background:#1e4d93; color:#fff; }}
    .wrap {{ padding:18px; }}
    table {{ width:100%; border-collapse:collapse; background:#fff; box-shadow:0 8px 22px rgba(15,35,75,.08); }}
    th {{ background:#1e4d93; color:#fff; padding:10px; }}
    td {{ padding:9px; text-align:center; border-bottom:1px solid #edf2fb; }}
    .car {{ font-weight:800; font-size:20px; letter-spacing:1px; }}
    .note {{ margin:12px 0 16px; color:#4b5d7a; }}
  </style>
</head>
<body>
  <div class="bar">
    <h1>نتيجة ترتيب الشيتات التي نُفذت هنا</h1>
    <p>7 جهات — {sum(len(g['cars']) for g in groups)} سيارة — أوراق DB_ جاهزة من الشهر 6 حتى الآن</p>
  </div>
  <div class="tabs">{tab_html}</div>
  <div class="wrap">
    <p class="note">كل رقم سيارة تحت معتمده. الشيت الحي لم يتغيّر بعد لأن جوجل يحتاج دخول حسابكم. هذا الملف هو نفس الهيكل الذي سيُكتب هناك.</p>
    <table>
      <thead><tr><th>المعتمد / الجهة</th><th>النوع</th><th>رقم السيارة</th><th>ورقة القاعدة</th></tr></thead>
      <tbody>{''.join(rows_html)}</tbody>
    </table>
  </div>
</body>
</html>
"""
html_path = OUT_DIR / "preview.html"
html_path.write_text(html, encoding="utf-8")
print(xlsx_path)
print(html_path)
print("groups", len(groups), "cars", sum(len(g["cars"]) for g in groups))
