#!/usr/bin/env python3
"""Build the live client-database workbook for Google Sheets import."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path("/home/ubuntu/Downloads/client-databases.xlsx")
REPO_OUT = Path("/workspace/artifacts/sheet-organization/قواعد_بيانات_العملاء.xlsx")
ARTIFACT = Path("/opt/cursor/artifacts/قواعد_بيانات_العملاء.xlsx")

CLIENTS = [
    {"name": "سجاد صويره", "kind": "معتمد", "db": "DB_سجاد_صويره", "cars": ["15871", "16414", "16416"]},
    {"name": "ابراهيم", "kind": "معتمد", "db": "DB_ابراهيم", "cars": ["31378", "22549", "36375", "32028"]},
    {"name": "رواد ريادة", "kind": "معتمد", "db": "DB_رواد_ريادة", "cars": ["22351", "31669", "28123"]},
    {"name": "قيصر وارد", "kind": "معتمد", "db": "DB_قيصر_وارد", "cars": ["29684", "33687", "29635", "36290"]},
    {"name": "علي صبار", "kind": "معتمد", "db": "DB_علي_صبار", "cars": ["12207", "31896", "27591", "36341", "23589", "35837", "22006", "30706", "27912", "23917", "24189"]},
    {"name": "قيصر شمري", "kind": "معتمد", "db": "DB_قيصر_شمري", "cars": ["20585", "27685", "20756", "20858"]},
    {"name": "شركة", "kind": "شركة", "db": "DB_شركة", "cars": ["22973", "24057", "24382", "25710", "30353", "29744", "29555", "13417", "27740"]},
    {"name": "شركة يونيغاز", "kind": "معتمد", "db": "DB_شركة_يونيغاز", "cars": ["22219", "22084", "22076", "22339"]},
]

LEDGER = [
    "رقم الوصل", "السائق", "رقم السيارة", "تاريخ التحميل", "تاريخ التفريغ",
    "الكمية طن", "الوجهة", "نوع الحركة", "الشهر", "الفترة", "سعر الطن",
    "المبلغ", "لترات الكاز", "قيمة الكاز", "المعتمد", "وقت الترحيل", "مفتاح الترحيل",
]
STATION = ["رقم الوصل", "اسم السائق", "رقم السيارة", "تاريخ التحميل", "تاريخ التفريغ", "الكمية", "المالك", "المحطة", "رابط الصورة", "وقت الإرسال", "لترات الكاز", "المسافة", "سعر النقل", "ملاحظات"]
FACTORY = ["رقم الوصل", "اسم السائق", "رقم السيارة", "تاريخ التحميل", "تاريخ التفريغ", "الكمية", "المالك", "المحطة/المعمل", "رابط الصورة", "وقت الإرسال", "مالك السيارة أو المالك", "ملاحظات", "source"]

HEAD_FILL = PatternFill("solid", fgColor="1E4D93")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="D0D7E2"),
    right=Side(style="thin", color="D0D7E2"),
    top=Side(style="thin", color="D0D7E2"),
    bottom=Side(style="thin", color="D0D7E2"),
)


def style_header(ws, cols):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    for col in range(1, cols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col)].width = 18


def write_sheet(wb, name, headers, rows=None, first=False):
    ws = wb.active if first else wb.create_sheet(name[:31])
    if first:
        ws.title = name[:31]
    ws.append(headers)
    for row in rows or []:
        ws.append(row)
    style_header(ws, len(headers))
    return ws


def main():
    wb = Workbook()
    client_rows = [[c["name"], c["kind"], "1", c["db"]] for c in CLIENTS]
    write_sheet(wb, "العملاء", ["الاسم", "النوع", "فعال", "قاعدة البيانات"], client_rows, first=True)

    car_rows = []
    index_rows = []
    for c in CLIENTS:
        for plate in c["cars"]:
            car_rows.append([plate, "", c["kind"], c["name"], "1", ""])
            index_rows.append([c["name"], c["kind"], plate, c["db"]])
    write_sheet(wb, "السيارات", ["رقم السيارة", "السائق الافتراضي", "تابع لـ", "اسم العميل", "فعال", "ملاحظات"], car_rows)

    write_sheet(wb, "تنظيم_المعتمدين", ["العميل", "النوع", "رقم السيارة", "ورقة القاعدة"], index_rows)

    for c in CLIENTS:
        write_sheet(wb, c["db"], LEDGER)
    write_sheet(wb, "DB_غير_مصنف", LEDGER)
    write_sheet(wb, "2026_09", STATION)
    write_sheet(wb, "F_2026_09", FACTORY)

    for path in (OUT, REPO_OUT, ARTIFACT):
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        print("wrote", path, "sheets", len(wb.sheetnames))
    print("clients", len(CLIENTS), "cars", len(car_rows))
    print("tabs", wb.sheetnames)


if __name__ == "__main__":
    main()
