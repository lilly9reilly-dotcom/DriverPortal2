#!/usr/bin/env python3
"""Organize the downloaded live Google Sheet by registered agent plates."""
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SEED = json.loads(Path("/tmp/org-seed.json").read_text(encoding="utf-8"))
GROUPS = SEED["groups"]
HEADERS = SEED["headers"]
SRC = Path("/tmp/sheet.xlsx")
OUT = Path("/opt/cursor/artifacts/sheet-organization/تنظيم_من_الشيت_الحي.xlsx")
HTML = Path("/opt/cursor/artifacts/sheet-organization/live-preview.html")
SUMMARY = Path("/opt/cursor/artifacts/sheet-organization/live-summary.json")

STATION = 34300
FACTORY = 8500
GAS = 430
MONTH_SHEETS = [
    "2026_06", "F_2026_06",
    "2026_07", "F_2026_07",
    "2026_08", "F_2026_08",
    "2026_09", "F_2026_09",
]


def ar_digits(s):
    return str(s or "").translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))


def plate_digits(v):
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, float):
        v = str(int(v)) if v == int(v) else str(v)
    elif isinstance(v, int):
        v = str(v)
    s = ar_digits(v).strip()
    s = re.sub(r"\.0+$", "", s)
    return re.sub(r"\D", "", s)


def digits_only(v):
    return plate_digits(v)


def to_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = ar_digits(v).replace(",", "").replace(" ", "")
    s = re.sub(r"[^\d.\-]", "", s)
    try:
        return float(s) if s else 0.0
    except ValueError:
        return 0.0


def qty_ton(v):
    n = to_num(v)
    if n <= 0 or n > 1000000:
        return 0.0
    if n >= 1000:
        n = n / 1000.0
    return round(n + 1e-12, 3)


def fmt_date(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v or "").strip()[:32]


def build_fleet():
    by_plate = {}
    for g in GROUPS:
        for car in g["cars"]:
            plate = digits_only(car["carNumber"])
            by_plate[plate] = g
    return by_plate


def match_group(car, fleet):
    key = plate_digits(car)
    if not key:
        return None, ""
    if key in fleet:
        return fleet[key], key
    best = None
    best_plate = ""
    best_len = 0
    for plate in fleet:
        if len(plate) < 4:
            continue
        hit = key == plate
        if not hit and len(key) >= len(plate):
            hit = key.endswith(plate)
        if not hit and len(plate) > len(key) >= 4:
            hit = plate.endswith(key)
        if hit and min(len(key), len(plate)) > best_len:
            best_len = min(len(key), len(plate))
            best = fleet[plate]
            best_plate = plate
    return best, best_plate


def period_label(date_value):
    text = fmt_date(date_value)
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", text)
    if not m:
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
        if m:
            day = int(m.group(1) if int(m.group(1)) <= 31 else m.group(2))
            return "1-15" if day <= 15 else "16-نهاية الشهر"
        return "كامل الشهر"
    day = int(m.group(3))
    return "1-15" if day <= 15 else "16-نهاية الشهر"


def header_map(row):
    m = {}
    for i, name in enumerate(row or []):
        token = re.sub(r"\s+", "", ar_digits(name or "")).replace("أ", "ا").replace("إ", "ا").replace("ة", "ه")
        if token and token not in m:
            m[token] = i
    return m


def cell(row, colmap, aliases, fallback):
    for a in aliases:
        if a in colmap and colmap[a] < len(row):
            return row[colmap[a]]
    if fallback is not None and fallback < len(row):
        return row[fallback]
    return ""


def read_receipts(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    fleet = build_fleet()
    out = []
    for sheet_name in MONTH_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        colmap = header_map(header)
        factory = sheet_name.upper().startswith("F_")
        month = sheet_name.replace("F_", "")
        for row in rows:
            if not row or not any(row):
                continue
            doc = cell(row, colmap, ["رقمالوصل"], 0)
            driver = cell(row, colmap, ["اسمالسائق", "السائق"], 1)
            car = cell(row, colmap, ["رقمالسيارة"], 2)
            if factory:
                qty = cell(row, colmap, ["الكميه", "الكمية"], 3 if "الكميه" in colmap or "الكمية" in colmap else 5)
                dest = cell(row, colmap, ["اسمالمعمل", "المحطه/المعمل", "المحطة/المعمل", "الجهه"], 4)
                load = cell(row, colmap, ["تاريخالتحميل"], 9 if sheet_name == "F_2026_06" else 3)
                unload = cell(row, colmap, ["تاريخالتفريغ"], 5 if sheet_name == "F_2026_06" else 4)
                liters = 0
            else:
                qty = cell(row, colmap, ["الكميه", "الكمية"], 5)
                dest = cell(row, colmap, ["المحطه", "المحطة", "الجهه"], 7)
                load = cell(row, colmap, ["تاريخالتحميل"], 3)
                unload = cell(row, colmap, ["تاريخالتفريغ"], 4)
                liters = to_num(cell(row, colmap, ["لتراتالكاز"], 10))
                if liters > 5000:
                    liters = 0
            owner = cell(row, colmap, ["المالك", "مالكالسيارهاوالمالك"], 6)
            if not str(doc or "").strip() and not str(car or "").strip():
                continue
            tons = qty_ton(qty)
            is_factory = factory or "معمل" in str(dest or "")
            price = FACTORY if is_factory else STATION
            amount = round(tons * price)
            gas = round(max(0, liters) * GAS)
            group, matched_plate = match_group(car, fleet)
            display_car = car
            if isinstance(display_car, float) and display_car == int(display_car):
                display_car = str(int(display_car))
            out.append({
                "doc": "" if doc is None else str(int(doc)) if isinstance(doc, float) and doc == int(doc) else str(doc),
                "driver": str(driver or "").strip(),
                "car": str(display_car or "").strip(),
                "car_key": matched_plate or plate_digits(car),
                "owner": str(owner or "").strip(),
                "load": fmt_date(load),
                "unload": fmt_date(unload),
                "period": period_label(unload or load),
                "qty": tons,
                "dest": str(dest or "").strip(),
                "kind": "معمل" if is_factory else "محطة",
                "month": month,
                "sheet": sheet_name,
                "price": price,
                "amount": amount,
                "liters": liters,
                "gas": gas,
                "agent": group["name"] if group else "غير مصنف",
                "db": group["dbSheet"] if group else "DB_غير_مصنف",
                "owner_kind": group["kind"] if group else "",
            })
    wb.close()
    return out


def write_xlsx(receipts):
    thin = Border(
        left=Side(style="thin", color="D0D7E2"),
        right=Side(style="thin", color="D0D7E2"),
        top=Side(style="thin", color="D0D7E2"),
        bottom=Side(style="thin", color="D0D7E2"),
    )
    head_fill = PatternFill("solid", fgColor="1E4D93")
    head_font = Font(bold=True, color="FFFFFF")
    wb = Workbook()

    by_agent = defaultdict(list)
    by_month = defaultdict(lambda: {"trips": 0, "qty": 0.0, "amount": 0, "unclassified": 0})
    for r in receipts:
        by_agent[r["db"]].append(r)
        m = by_month[r["month"]]
        m["trips"] += 1
        m["qty"] += r["qty"]
        m["amount"] += r["amount"]
        if r["agent"] == "غير مصنف":
            m["unclassified"] += 1

    ws = wb.active
    ws.title = "تنظيم_المعتمدين"
    ws.sheet_view.rightToLeft = True
    ws.append(["ترتيب الشيت الحي — من 2026_06 حتى 2026_09"])
    ws.append(["إجمالي الوصولات المفروزة", len(receipts)])
    ws.append(["مصنفة", sum(1 for r in receipts if r["agent"] != "غير مصنف")])
    ws.append(["غير مصنفة", sum(1 for r in receipts if r["agent"] == "غير مصنف")])
    ws.append([])
    ws.append(["المعتمد / الجهة", "النوع", "رقم السيارة", "ورقة القاعدة", "وصولات هذه السيارة", "كمية (طن)", "المبلغ"])
    for cell in ws[6]:
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center")

    car_stats = defaultdict(lambda: {"trips": 0, "qty": 0.0, "amount": 0, "agent": "", "kind": "", "db": ""})
    for r in receipts:
        if r["agent"] == "غير مصنف":
            continue
        key = (r["agent"], r["car_key"] or r["car"])
        s = car_stats[key]
        s["trips"] += 1
        s["qty"] += r["qty"]
        s["amount"] += r["amount"]
        s["agent"] = r["agent"]
        s["kind"] = r["owner_kind"]
        s["db"] = r["db"]

    # keep registered cars even if 0 trips
    fleet = build_fleet()
    seen = set()
    for g in GROUPS:
        for car in g["cars"]:
            plate = digits_only(car["carNumber"])
            seen.add((g["name"], plate))
            s = car_stats[(g["name"], plate)]
            s["agent"] = g["name"]
            s["kind"] = g["kind"]
            s["db"] = g["dbSheet"]
            ws.append([g["name"], g["kind"], car["carNumber"], g["dbSheet"], s["trips"], round(s["qty"], 3), s["amount"]])

    ws.append([])
    ws.append(["ملخص الشهر", "وصولات", "غير مصنف", "كمية", "المبلغ"])
    for month in sorted(by_month):
        m = by_month[month]
        ws.append([month, m["trips"], m["unclassified"], round(m["qty"], 3), m["amount"]])

    ws.append([])
    ws.append(["ملخص الجهة", "وصولات", "كمية", "المبلغ"])
    for g in GROUPS:
        rows = by_agent[g["dbSheet"]]
        ws.append([g["name"], len(rows), round(sum(x["qty"] for x in rows), 3), sum(x["amount"] for x in rows)])
    rows = by_agent["DB_غير_مصنف"]
    ws.append(["غير مصنف", len(rows), round(sum(x["qty"] for x in rows), 3), sum(x["amount"] for x in rows)])

    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 22

    ws = wb.create_sheet("Agents")
    ws.sheet_view.rightToLeft = True
    ws.append(["الاسم", "النوع", "فعال", "ملاحظات"])
    for g in GROUPS:
        ws.append([g["name"], g["kind"], "1", f"وصولات مفروزة: {len(by_agent[g['dbSheet']])}"])

    ws = wb.create_sheet("Fleet")
    ws.sheet_view.rightToLeft = True
    ws.append(["رقم السيارة", "السائق الافتراضي", "تابع لـ", "اسم المعتمد", "فعال", "ملاحظات"])
    for g in GROUPS:
        for car in g["cars"]:
            ws.append([car["carNumber"], "", g["kind"], g["name"], "1", ""])

    ledger_headers = HEADERS

    def write_db(name, rows):
        ws = wb.create_sheet(name[:31])
        ws.sheet_view.rightToLeft = True
        ws.append(ledger_headers)
        for r in rows:
            ws.append([
                r["doc"], r["driver"], r["car"], r["load"], r["unload"], r["qty"], r["dest"],
                r["kind"], r["month"], r.get("period") or period_label(r["unload"] or r["load"]),
                r["price"], r["amount"], r["liters"], r["gas"], r["agent"], "",
                f"{r['doc']}|{r['car_key']}|{r['load']}|{r['kind']}"
            ])
        for col in range(1, 18):
            ws.column_dimensions[get_column_letter(col)].width = 16

    for g in GROUPS:
        write_db(g["dbSheet"], by_agent[g["dbSheet"]])
    write_db("DB_غير_مصنف", by_agent["DB_غير_مصنف"])

    unknown = by_agent["DB_غير_مصنف"]
    unknown_cars = defaultdict(lambda: {"trips": 0, "qty": 0.0, "amount": 0, "owners": set(), "samples": []})
    for r in unknown:
        s = unknown_cars[r["car"] or r["car_key"] or "?"]
        s["trips"] += 1
        s["qty"] += r["qty"]
        s["amount"] += r["amount"]
        if r.get("owner"):
            s["owners"].add(r["owner"])
        if len(s["samples"]) < 2:
            s["samples"].append(r.get("driver") or "")

    ws = wb.create_sheet("سيارات_غير_مسجلة")
    ws.sheet_view.rightToLeft = True
    ws.append(["رقم السيارة", "وصولات", "كمية طن", "المبلغ", "المالك كما كتبه السائق", "سائقون"])
    for cell in ws[1]:
        cell.fill = head_fill
        cell.font = head_font
    for car, s in sorted(unknown_cars.items(), key=lambda item: -item[1]["trips"]):
        ws.append([
            car, s["trips"], round(s["qty"], 3), s["amount"],
            " / ".join(sorted(s["owners"]))[:80],
            " / ".join(x for x in s["samples"] if x)[:80],
        ])
    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 24

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return by_agent, by_month, unknown_cars


def write_html(receipts, by_agent, by_month, unknown_cars):
    classified = [r for r in receipts if r["agent"] != "غير مصنف"]
    rows = []
    for g in GROUPS:
        recs = by_agent[g["dbSheet"]]
        rows.append(
            f"<tr><td>{g['name']}</td><td>{g['kind']}</td><td>{len(g['cars'])}</td>"
            f"<td>{len(recs)}</td><td>{round(sum(x['qty'] for x in recs),3)}</td>"
            f"<td>{sum(x['amount'] for x in recs):,}</td><td>{g['dbSheet']}</td></tr>"
        )
    months = "".join(
        f"<tr><td>{m}</td><td>{by_month[m]['trips']}</td><td>{by_month[m]['unclassified']}</td>"
        f"<td>{round(by_month[m]['qty'],3)}</td><td>{by_month[m]['amount']:,}</td></tr>"
        for m in sorted(by_month)
    )
    unknown_rows = "".join(
        f"<tr><td>{car}</td><td>{s['trips']}</td><td>{round(s['qty'],3)}</td>"
        f"<td>{s['amount']:,}</td><td>{' / '.join(sorted(s['owners']))}</td></tr>"
        for car, s in sorted(unknown_cars.items(), key=lambda item: -item[1]["trips"])[:20]
    )
    html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head><meta charset="UTF-8">
<title>فرز الشيت الحي</title>
<link href="https://fonts.googleapis.com/css2?family=Tajawal:wght@500;800&display=swap" rel="stylesheet">
<style>
body{{margin:0;font-family:Tajawal,sans-serif;background:#eef2f7;color:#17253f}}
.bar{{background:#0b3a74;color:#fff;padding:18px 22px}}
.bar h1{{margin:0;font-size:28px}}
.wrap{{padding:18px}}
.cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin-bottom:16px}}
.card{{background:#fff;border-radius:12px;padding:12px;box-shadow:0 8px 22px rgba(15,35,75,.08)}}
.card b{{display:block;font-size:28px;color:#1e4d93}}
table{{width:100%;border-collapse:collapse;background:#fff;margin-bottom:18px}}
th{{background:#1e4d93;color:#fff;padding:8px}} td{{padding:8px;text-align:center;border-bottom:1px solid #edf2fb}}
.note{{background:#fff7e6;border:1px solid #f0d48a;padding:12px;border-radius:10px;margin-bottom:16px}}
</style></head>
<body>
<div class="bar"><h1>فرز الشيت الحي — من يونيو حتى سبتمبر 2026</h1>
<p>المصدر: Google Sheet 1adlJxYSgBftTcagTAyl9GrwNwrBTPOrTDQbYeBhqCiM</p></div>
<div class="wrap">
<div class="note">الأشهر الأصلية لم تُحذف. هذه أوراق جديدة: Agents وFleet وDB لكل معتمد، بالإضافة إلى سيارات غير مسجّلة.</div>
<div class="cards">
<div class="card">إجمالي الوصولات<b>{len(receipts):,}</b></div>
<div class="card">مصنفة حسب السيارة<b>{len(classified):,}</b></div>
<div class="card">غير مصنفة<b>{len(receipts)-len(classified):,}</b></div>
<div class="card">جهات مسجّلة<b>{len(GROUPS)}</b></div>
</div>
<h2>المعتمدون بعد الفرز</h2>
<table><thead><tr><th>الجهة</th><th>النوع</th><th>سيارات مسجّلة</th><th>وصولات</th><th>كمية طن</th><th>المبلغ</th><th>الورقة</th></tr></thead>
<tbody>{''.join(rows)}</tbody></table>
<h2>حسب الشهر</h2>
<table><thead><tr><th>الشهر</th><th>وصولات</th><th>غير مصنف</th><th>كمية</th><th>المبلغ</th></tr></thead>
<tbody>{months}</tbody></table>
<h2>أكثر السيارات غير المسجّلة</h2>
<table><thead><tr><th>السيارة</th><th>وصولات</th><th>كمية</th><th>المبلغ</th><th>المالك في الشيت</th></tr></thead>
<tbody>{unknown_rows}</tbody></table>
</div></body></html>"""
    HTML.write_text(html, encoding="utf-8")


def main():
    receipts = read_receipts(SRC)
    by_agent, by_month, unknown_cars = write_xlsx(receipts)
    write_html(receipts, by_agent, by_month, unknown_cars)
    summary = {
        "total": len(receipts),
        "classified": sum(1 for r in receipts if r["agent"] != "غير مصنف"),
        "unclassified": sum(1 for r in receipts if r["agent"] == "غير مصنف"),
        "byAgent": {g["name"]: len(by_agent[g["dbSheet"]]) for g in GROUPS},
        "byMonth": {m: by_month[m]["trips"] for m in by_month},
        "unknownCars": len(unknown_cars),
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print("xlsx", OUT)
    print("html", HTML)


if __name__ == "__main__":
    main()
